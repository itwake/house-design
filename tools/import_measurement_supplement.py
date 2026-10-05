"""Read the R2 survey workbook; never edit/recalculate the supplied workbook.

Usage: python tools/import_measurement_supplement.py path/to/实测尺寸核对表_补充复核.xlsx
Cached formula values and formula text are preserved separately. Blank adoption
cells stay blank, including when the calculation column contains a number.
"""
from collections import Counter
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PREVIOUS = ROOT / "models/measurements-20261004.json"
DESTINATION = ROOT / "models/measurements-20261004-r2.json"

RECORD_FIELDS = [
    "id", "room", "label", "rawText", "rawNumber", "rawUnit",
    "convertedOrCalculatedMm", "adoptedMm", "status",
    "confidenceAndReference", "issueIds", "unused", "sourceNote",
    "effectiveCandidates", "note",
]
ISSUE_FIELDS = [
    "id", "room", "problem", "check", "priority", "confirmedMm",
    "explanation", "status", "sourceNote", "reviewNote",
]
# The sheet's A3 summary and C/D/J descriptions classify U01 as resolved even
# though its H6 legacy formula cache says "已填写，待复核". Preserve both meanings.
REPORT_GROUPS = {
    "已解决": ["U01", "Q02", "Q03", "Q06", "Q08", "Q09", "Q10", "Q13",
               "Q16", "Q17", "Q18", "Q19"],
    "用户接受偏差": ["Q11", "Q12"],
    "未完成": ["Q01", "Q04", "Q05", "Q07", "Q14", "Q15", "Q20", "Q21"],
}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalized_number(value):
    """Suppress binary float noise only; exact workbook cache is kept separately."""
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        value = round(value, 5)
        return int(value) if value == int(value) else value
    return value


def cell_evidence(sheet_name, row_number, row, formula_row, fields):
    cells = {}
    formulas = {}
    for column, field in enumerate(fields, 1):
        if field == "unused":
            continue
        address = f"{sheet_name}!{get_column_letter(column)}{row_number}"
        cells[field] = address
        formula = formula_row[column - 1] if column <= len(formula_row) else None
        if isinstance(formula, str) and formula.startswith("="):
            formulas[field] = {
                "cell": address,
                "formula": formula,
                "cachedValue": row[column - 1] if column <= len(row) else None,
            }
    return cells, formulas


def read_rows(values_book, formula_book, sheet_name, fields):
    values = list(values_book[sheet_name].values)
    formulas = list(formula_book[sheet_name].values)
    assert len(values) == len(formulas), f"{sheet_name}: formula/cache row mismatch"
    result = []
    for row_number, (row, formula_row) in enumerate(zip(values, formulas), 1):
        if row_number <= 5 or not row[0]:
            continue
        record = {field: row[index] if index < len(row) else None
                  for index, field in enumerate(fields) if field != "unused"}
        record["sourceCells"], record["formulaCells"] = cell_evidence(
            sheet_name, row_number, row, formula_row, fields)
        result.append(record)
    return result


def record_diff(previous, records):
    old = {record["id"]: record for record in previous["records"]}
    current = {record["id"]: record for record in records}
    changes = []
    calculation_changes = []
    for record in records:
        prior = old.get(record["id"])
        if prior is None:
            continue
        for field, destination in [("adoptedMm", changes),
                                   ("convertedOrCalculatedMm", calculation_changes)]:
            before = normalized_number(prior.get(field))
            after = record.get(field)
            if before == after:
                continue
            destination.append({
                "id": record["id"], "room": record["room"],
                "label": record["label"], "previous": before, "current": after,
                "previousCell": prior.get("cell") if field == "adoptedMm" else
                    prior.get("cell", "").replace("!I", "!H"),
                "currentCell": record["sourceCells"][field],
                "adoptionType": record["status"], "note": record["note"],
            })
    return {
        "scope": "只比较上一版原表采用值和换算值；不引入旧图估值。历史采用值仅供版本追溯，不是当前候选。",
        "adoptedValueChanges": changes,
        "calculationValueChanges": calculation_changes,
        "addedRecordIds": [key for key in current if key not in old],
        "removedRecordIds": [key for key in old if key not in current],
    }


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    source = Path(sys.argv[1]).resolve()
    before_sha = sha256(source)
    previous = json.loads(PREVIOUS.read_text(encoding="utf-8"))
    values_book = load_workbook(source, data_only=True, read_only=True)
    formula_book = load_workbook(source, data_only=False, read_only=True)
    try:
        headers = list(values_book["尺寸核对"].values)[4]
        assert headers[7] == "本轮采用\nmm" and headers[8] == "采用类型"
        assert headers[3] == "原始记录" and headers[14] == "本轮复核说明"
        records = read_rows(values_book, formula_book, "尺寸核对", RECORD_FIELDS)
        for record in records:
            record["cell"] = record["sourceCells"]["adoptedMm"]
            record["rawCell"] = record["sourceCells"]["rawText"]
            record["adoptionType"] = record["status"]
            for field in ("convertedOrCalculatedMm", "adoptedMm"):
                record[field] = normalized_number(record[field])
            for field, alias in [("convertedOrCalculatedMm", "calculationFormula"),
                                 ("adoptedMm", "adoptionFormula")]:
                if field in record["formulaCells"]:
                    record[alias] = record["formulaCells"][field]["formula"]
        issues = read_rows(values_book, formula_book, "待核清单", ISSUE_FIELDS)
        for issue in issues:
            row_number = issue["sourceCells"]["id"].split("!A")[1]
            issue["cell"] = f"待核清单!A{row_number}:J{row_number}"
            issue["reportClassification"] = next(
                group for group, ids in REPORT_GROUPS.items() if issue["id"] in ids)
            issue["classificationSource"] = (
                f"待核清单!A3及C{row_number}:J{row_number}复核说明；不覆盖H列原状态")
        area_rows = list(values_book["区域汇总"].values)
        areas = [dict(zip(["room", "summary", "pending", "drawingStatus", "sourceNote"], row[:5]),
                      cell=f"区域汇总!A{index}:E{index}")
                 for index, row in enumerate(area_rows, 1) if index > 5 and row[0]]
        summary_note = list(values_book["待核清单"].values)[2][0]
    finally:
        values_book.close()
        formula_book.close()

    assert len(records) == 158 and len({r["id"] for r in records}) == 158
    assert len(issues) == 22 and len({i["id"] for i in issues}) == 22
    assert {i["id"] for i in issues} == {i for ids in REPORT_GROUPS.values() for i in ids}
    assert sha256(source) == before_sha, "Source workbook changed during read-only import"
    diff = record_diff(previous, records)
    result = {
        "version": "2.0", "date": "2026-10-04", "revision": "r2", "units": "mm",
        "sourceName": source.name, "sourceSha256": before_sha,
        "sourceRole": "用户提供的补充复尺核对表；记录确认、引用、推算和待核，不等同施工图或许可",
        "adoptionRule": "本版本H列为本轮采用值，G列为换算/计算，I列为采用类型。采用留空不得用G列、原图或上一版旧估值回填。数值不自动证明绝对定位或测量参考面已确定。",
        "formulaHandling": "只读保存文件已有公式缓存；未计算或改写工作簿。formulaCells同时保留公式原文与未经数值归整的缓存；数值字段仅归整二进制小数尾差。",
        "historicalValueRule": "不从上一版补入旧图估值或旧候选。新版N列原文按源表保留，但仍须结合I列采用类型，已划除及待核项不得当作已采用。被更正的照片读数仅在原始转录和明确的版本差异中保留，不作为当前采用值。",
        "orientation": previous["orientation"],
        "roomMapping": previous["roomMapping"],
        "roomMappingBasis": previous["roomMappingBasis"],
        "previousSource": {
            "path": "models/measurements-20261004.json",
            "sha256": sha256(PREVIOUS), "sourceName": previous["sourceName"],
            "sourceSha256": previous["sourceSha256"],
        },
        "summary": {
            "records": len(records), "issues": len(issues),
            "adoptionTypes": dict(Counter(r["status"] for r in records)),
            "rawIssueStatuses": dict(Counter(i["status"] for i in issues)),
            "reportClassifications": {key: len(ids) for key, ids in REPORT_GROUPS.items()},
            "classificationSource": "待核清单!A3及各行C/D/J本轮复核说明",
            "classificationNote": "报告分类与H列原处理状态分开保存。尤其U01的H列缓存仍为已填写，待复核，不能把22项或所有有数值记录宣称全部确认。",
            "sourceSummaryText": summary_note,
            "adoptedValueChangeCount": len(diff["adoptedValueChanges"]),
            "newRecordCount": len(diff["addedRecordIds"]),
        },
        "areas": areas, "records": records, "issues": issues, "diff": diff,
    }
    DESTINATION.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"destination": str(DESTINATION), "sourceSha256": before_sha,
                      **result["summary"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
