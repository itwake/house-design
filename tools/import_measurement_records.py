"""Read the supplied workbook without modifying it; preserve adopted vs raw values.

Usage: python tools/import_measurement_records.py path/to/实测尺寸核对表.xlsx
The public JSON contains dimension records only, not the user's original documents.
"""
from pathlib import Path
import hashlib
import json
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
source = Path(sys.argv[1])
book = load_workbook(source, data_only=True, read_only=True)
formulas = load_workbook(source, data_only=False, read_only=True)
records = []
for row_number, row in enumerate(book['尺寸核对'].values, 1):
    if row_number <= 5 or not row[0]:
        continue
    r = dict(zip(['id', 'room', 'label', 'oldEstimate', 'rawText', 'rawNumber',
                  'rawUnit', 'convertedOrCalculatedMm', 'adoptedMm', 'status',
                  'differenceMm', 'confidenceAndReference', 'issueIds',
                  'unused', 'sourceNote', 'historicalCandidates', 'note'], row))
    r.pop('unused', None)
    r['cell'] = f'尺寸核对!I{row_number}'
    r['rawCell'] = f'尺寸核对!E{row_number}'
    formula = formulas['尺寸核对'].cell(row_number, 8).value
    if isinstance(formula, str) and formula.startswith('='):
        r['calculationFormula'] = formula
    for key in ['convertedOrCalculatedMm', 'adoptedMm', 'differenceMm']:
        if isinstance(r[key], (int, float)):
            r[key] = round(r[key], 5)
    records.append(r)
issues = []
for row_number, row in enumerate(book['待核清单'].values, 1):
    if row_number <= 5 or not row[0]:
        continue
    issue = dict(zip(['id', 'room', 'problem', 'check', 'priority',
                      'confirmedMm', 'explanation', 'status', 'sourceNote'], row))
    issue['cell'] = f'待核清单!A{row_number}:I{row_number}'
    issues.append(issue)
result = {
    'version': '1.0', 'date': '2026-10-04', 'units': 'mm',
    'sourceName': source.name,
    'sourceSha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'sourceRole': '用户提供的现状复尺核对表；不是改造后设计图或施工许可',
    'adoptionRule': 'I列为原表本轮采用值，H列可能只是换算/计算；有采用数值也不代表测量端面、空间定位或全屋尺寸链已确定。',
    'orientation': '九张分房图逆时针旋转90°对应网站北向；原分房图上下不代表实地正北。',
    'roomMapping': {'卧室1': 'room_c', '卧室2': 'room_b', '卧室3': 'room_a'},
    'roomMappingBasis': '依据窗位、缺口和卧室3套卫相邻关系判读；保留原表区域名称供核对。',
    'records': records, 'issues': issues,
}
destination = ROOT / 'models/measurements-20261004.json'
destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps({'records': len(records), 'issues': len(issues), 'destination': str(destination)}, ensure_ascii=False))
